import csv
import hashlib
import io
import json
import mimetypes
import re
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
STATIC = ROOT / "static"
DB_PATH = Path.home() / "OneDrive" / "Desktop" / "pharmacy-analytics-dashboard" / "pharmacy.db"
HOST = "127.0.0.1"
PORT = 8001
MAX_UPLOAD_SIZE = 50 * 1024 * 1024
HOSPITALS = ("Surya Hospital", "NIMMS Hospital (Haldoni)", "Vedansh Hospital (Dadri)")
HOSPITAL_PHARMACIES = {
    "Surya Hospital": ("IPD", "OPD"),
    "NIMMS Hospital (Haldoni)": ("Pharmacy",),
    "Vedansh Hospital (Dadri)": ("Pharmacy",),
}
PHARMACIES = tuple(dict.fromkeys(
    pharmacy for pharmacies in HOSPITAL_PHARMACIES.values() for pharmacy in pharmacies
))
LEGACY_HOSPITALS = {
    "Hospital 1": "Surya Hospital",
    "Hospital 2": "NIMMS Hospital (Haldoni)",
    "Hospital 3": "Vedansh Hospital (Dadri)",
}
SINGLE_PHARMACY_HOSPITALS = HOSPITALS[1:]
LEGACY_SINGLE_PHARMACY_HOSPITALS = tuple(
    legacy for legacy, current in LEGACY_HOSPITALS.items() if current in SINGLE_PHARMACY_HOSPITALS
)
HOSPITAL_SQL = "CASE u.hospital " + " ".join(
    "WHEN '{}' THEN '{}'".format(legacy.replace("'", "''"), current.replace("'", "''"))
    for legacy, current in LEGACY_HOSPITALS.items()
) + " ELSE u.hospital END"
PHARMACY_SQL = (
    "CASE WHEN u.hospital IN ({}) AND u.pharmacy IN ('IPD', 'OPD') "
    "THEN 'Pharmacy' ELSE u.pharmacy END"
).format(", ".join("'{}'".format(name.replace("'", "''")) for name in (
    *SINGLE_PHARMACY_HOSPITALS, *LEGACY_SINGLE_PHARMACY_HOSPITALS
)))
REPORT_TYPES = ("sales", "collection", "purchase")
REPORT_LABELS = {
    "sales": "Sales Revenue",
    "collection": "Collection Revenue",
    "purchase": "Purchase Value",
}
FIELDS = (
    "date", "amount", "product", "supplier", "quantity", "unit_cost",
    "reference", "cash_due", "credit_due",
)
FIELD_LABELS = {
    "date": "Transaction date",
    "amount": "Line amount / net amount",
    "product": "Product / SKU",
    "supplier": "Supplier name",
    "quantity": "Quantity",
    "unit_cost": "Unit rate / purchase cost",
    "reference": "Bill / transaction reference",
    "cash_due": "Cash due",
    "credit_due": "Credit due",
}
ALIASES = {
    "date": ("bill date", "transaction date", "invoice date", "voucher date", "date"),
    "amount": ("final amt", "final amount", "net amount", "net amt", "net sale", "sales amount", "revenue", "amount", "amt", "total", "net value", "line total"),
    "product": ("product", "product name", "medicine", "medicine name", "item", "item name", "sku", "item code"),
    "supplier": ("supplier name", "supplier", "vendor name", "vendor", "party name", "party", "distributor name", "distributor"),
    "quantity": ("quantity", "qty", "sold qty", "purchase qty", "units"),
    "unit_cost": ("rate", "unit rate", "purchase rate", "cost", "unit cost", "purchase price"),
    "reference": ("bill no", "bill number", "invoice no", "invoice number", "voucher no", "reference", "transaction id"),
    "cash_due": ("cash due", "cash due amount"),
    "credit_due": ("credit due", "credit due amount"),
}


@contextmanager
def connect():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def init_db():
    with connect() as connection:
        connection.execute("PRAGMA journal_mode = WAL")
        connection.execute(
            """CREATE TABLE IF NOT EXISTS uploads (
                id INTEGER PRIMARY KEY,
                hospital TEXT NOT NULL,
                pharmacy TEXT NOT NULL,
                report_type TEXT NOT NULL,
                filename TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                row_count INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE (hospital, pharmacy, report_type, sha256)
            )"""
        )
        upload_columns = {
            row["name"] for row in connection.execute("PRAGMA table_info(uploads)")
        }
        if "source_file" not in upload_columns:
            connection.execute("ALTER TABLE uploads ADD COLUMN source_file BLOB")
        if "skipped_count" not in upload_columns:
            connection.execute(
                "ALTER TABLE uploads ADD COLUMN skipped_count INTEGER NOT NULL DEFAULT 0"
            )
        if "due_data_loaded" not in upload_columns:
            connection.execute(
                "ALTER TABLE uploads ADD COLUMN due_data_loaded INTEGER NOT NULL DEFAULT 0"
            )
        if "supplier_data_loaded" not in upload_columns:
            connection.execute(
                "ALTER TABLE uploads ADD COLUMN supplier_data_loaded INTEGER NOT NULL DEFAULT 0"
            )
        connection.execute(
            """CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY,
                upload_id INTEGER NOT NULL REFERENCES uploads(id) ON DELETE CASCADE,
                row_number INTEGER NOT NULL,
                txn_date TEXT NOT NULL,
                product TEXT NOT NULL DEFAULT '',
                supplier TEXT NOT NULL DEFAULT '',
                quantity REAL,
                amount REAL NOT NULL,
                unit_cost REAL,
                reference TEXT NOT NULL DEFAULT '',
                UNIQUE (upload_id, row_number)
            )"""
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_transactions_date ON transactions(txn_date)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_transactions_upload ON transactions(upload_id)"
        )
        transaction_columns = {
            row["name"] for row in connection.execute("PRAGMA table_info(transactions)")
        }
        if "cash_due" not in transaction_columns:
            connection.execute("ALTER TABLE transactions ADD COLUMN cash_due REAL")
        if "credit_due" not in transaction_columns:
            connection.execute("ALTER TABLE transactions ADD COLUMN credit_due REAL")
        if "supplier" not in transaction_columns:
            connection.execute(
                "ALTER TABLE transactions ADD COLUMN supplier TEXT NOT NULL DEFAULT ''"
            )
    backfill_due_data()
    backfill_supplier_data()


def decode_cell(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if value is None:
        return ""
    return str(value).strip()


def read_rows(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        raw_rows = list(csv.reader(io.StringIO(text), dialect))
    elif suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise ValueError("Excel file nahi khul saki; file format check karein.") from error
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    elif suffix == ".xls":
        raise ValueError("Purana .xls format supported nahi hai; file ko .xlsx ya .csv mein save karein.")
    else:
        raise ValueError("Sirf .xlsx aur .csv files supported hain.")

    header_index = next(
        (index for index, row in enumerate(raw_rows) if any(decode_cell(cell) for cell in row)),
        None,
    )
    if header_index is None:
        raise ValueError("File mein header ya data nahi mila.")
    headers = []
    for index, cell in enumerate(raw_rows[header_index]):
        header = decode_cell(cell) or "Column {}".format(index + 1)
        if header in headers:
            header = "{} ({})".format(header, index + 1)
        headers.append(header)
    if len(headers) < 2:
        raise ValueError("Report mein kam se kam do columns hone chahiye.")
    rows = []
    for raw in raw_rows[header_index + 1:]:
        cells = list(raw) + [None] * max(0, len(headers) - len(raw))
        if not any(decode_cell(cell) for cell in cells):
            continue
        rows.append({header: decode_cell(cells[index]) for index, header in enumerate(headers)})
    return headers, rows


def parse_date(value):
    text = decode_cell(value)
    if not text:
        raise ValueError("date missing")
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        pass
    if re.fullmatch(r"\d{5,6}(?:\.0+)?", text):
        try:
            serial = float(text)
            return (datetime(1899, 12, 30) + timedelta(days=serial)).date().isoformat()
        except (OverflowError, ValueError):
            pass
    date_formats = (
        "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d",
        "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
        "%m/%d/%Y", "%m-%d-%Y", "%m.%d.%Y",
        "%d/%m/%y", "%d-%m-%y", "%d.%m.%y",
        "%d-%b-%Y", "%d %b %Y", "%d %B %Y",
        "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y",
    )
    formats = tuple(
        fmt
        for date_fmt in date_formats
        for fmt in (
            date_fmt,
            date_fmt + " %H:%M",
            date_fmt + " %H:%M:%S",
            date_fmt + " %H:%M:%S.%f",
            date_fmt + "T%H:%M",
            date_fmt + "T%H:%M:%S",
            date_fmt + "T%H:%M:%S.%f",
        )
    )
    for fmt in formats:
        try:
            return datetime.strptime(text.strip(), fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError("date format samajh nahi aaya ({})".format(text[:32]))


def parse_number(value):
    text = decode_cell(value).replace(",", "").replace("₹", "").replace("INR", "").replace("Rs.", "").strip()
    if not text:
        return None
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1].strip()
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text):
        return None
    return float(text)


def suggested_mapping(headers):
    normalized = {header: re.sub(r"[^a-z0-9]+", " ", header.lower()).strip() for header in headers}
    result = {}
    for field, aliases in ALIASES.items():
        exact = next(
            (header for alias in aliases for header, clean in normalized.items() if clean == alias),
            "",
        )
        result[field] = exact or next(
            (
                header for alias in sorted(aliases, key=len, reverse=True)
                for header, clean in normalized.items()
                if len(alias) > 3 and alias in clean
            ),
            "",
        )
    return result


def parse_multipart(handler):
    content_type = handler.headers.get("Content-Type", "")
    if not content_type.lower().startswith("multipart/form-data"):
        raise ValueError("Multipart file upload required hai.")
    try:
        length = int(handler.headers.get("Content-Length", "0"))
    except ValueError as error:
        raise ValueError("Invalid content length.") from error
    if length <= 0 or length > MAX_UPLOAD_SIZE:
        raise ValueError("Upload khaali hai ya 50 MB limit se badi hai.")
    raw = (
        "Content-Type: {}\r\nMIME-Version: 1.0\r\n\r\n".format(content_type).encode("latin-1")
        + handler.rfile.read(length)
    )
    message = BytesParser(policy=policy.default).parsebytes(raw)
    fields = {}
    files = []
    for part in message.walk():
        if part.is_multipart():
            continue
        name = part.get_param("name", header="content-disposition")
        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""
        if filename:
            files.append({
                "filename": Path(filename.replace("\\", "/")).name,
                "content": payload,
            })
        elif name:
            fields[name] = payload.decode("utf-8", errors="replace")
    if not files:
        raise ValueError("Excel/CSV file select karein.")
    return files, fields


def selected_filters(query):
    hospital = query.get("hospital", ["all"])[0]
    pharmacy = query.get("pharmacy", ["all"])[0]
    if hospital != "all" and hospital not in HOSPITALS:
        raise ValueError("Invalid hospital filter.")
    if pharmacy != "all" and pharmacy not in PHARMACIES:
        raise ValueError("Invalid pharmacy filter.")
    if hospital != "all" and pharmacy != "all" and pharmacy not in HOSPITAL_PHARMACIES[hospital]:
        raise ValueError("Selected pharmacy does not belong to this hospital.")
    return hospital, pharmacy


def hospital_values(hospital):
    return (hospital,) + tuple(
        legacy for legacy, current in LEGACY_HOSPITALS.items() if current == hospital
    )


def add_facility_filters(conditions, params, hospital, pharmacy):
    if hospital != "all":
        values = hospital_values(hospital)
        conditions.append("u.hospital IN ({})".format(",".join("?" for _ in values)))
        params.extend(values)
    if pharmacy == "Pharmacy":
        legacy_single = (*SINGLE_PHARMACY_HOSPITALS, *LEGACY_SINGLE_PHARMACY_HOSPITALS)
        conditions.append(
            "(u.pharmacy = ? OR (u.pharmacy IN ('IPD', 'OPD') AND u.hospital IN ({})))".format(
                ",".join("?" for _ in legacy_single)
            )
        )
        params.append("Pharmacy")
        params.extend(legacy_single)
    elif pharmacy != "all":
        conditions.append("u.pharmacy = ?")
        params.append(pharmacy)
        if pharmacy in ("IPD", "OPD"):
            surya_values = hospital_values(HOSPITALS[0])
            conditions.append("u.hospital IN ({})".format(",".join("?" for _ in surya_values)))
            params.extend(surya_values)


def duplicate_upload_exists(connection, hospital, pharmacy, report_type, digest):
    values = hospital_values(hospital)
    conditions = ["hospital IN ({})".format(",".join("?" for _ in values)), "report_type = ?", "sha256 = ?"]
    params = [*values, report_type, digest]
    if hospital in SINGLE_PHARMACY_HOSPITALS:
        conditions.append("pharmacy IN ('Pharmacy', 'IPD', 'OPD')")
    else:
        conditions.append("pharmacy = ?")
        params.append(pharmacy)
    return connection.execute(
        "SELECT 1 FROM uploads WHERE " + " AND ".join(conditions) + " LIMIT 1",
        params,
    ).fetchone() is not None


def prepare_import(filename, content, report_type, mapping):
    headers, rows = read_rows(filename, content)
    if report_type not in REPORT_TYPES:
        raise ValueError("Select a valid report type.")
    if not isinstance(mapping, dict):
        raise ValueError("Column mapping valid nahi hai.")
    if mapping.get("date") not in headers:
        raise ValueError("{} column map karna zaroori hai.".format(FIELD_LABELS["date"]))
    if report_type != "purchase" and mapping.get("amount") not in headers:
        raise ValueError("{} column map karna zaroori hai.".format(FIELD_LABELS["amount"]))
    if report_type == "purchase" and mapping.get("amount") not in headers:
        if mapping.get("quantity") not in headers or mapping.get("unit_cost") not in headers:
            raise ValueError("Purchase ke liye line amount ya quantity + unit cost map karein.")

    prepared = []
    due_by_row = {}
    supplier_by_row = {}
    skipped_count = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        try:
            txn_date = parse_date(row.get(mapping["date"]))
            quantity = parse_number(row.get(mapping.get("quantity", ""))) if mapping.get("quantity") else None
            amount = parse_number(row.get(mapping.get("amount", ""))) if mapping.get("amount") else None
            unit_cost = parse_number(row.get(mapping.get("unit_cost", ""))) if mapping.get("unit_cost") else None
            if amount is None and report_type == "purchase" and quantity is not None and unit_cost is not None:
                amount = quantity * unit_cost
            if amount is None:
                raise ValueError("amount missing or invalid")
            product = row.get(mapping.get("product", ""), "") if mapping.get("product") else ""
            supplier = row.get(mapping.get("supplier", ""), "") if mapping.get("supplier") else ""
            reference = row.get(mapping.get("reference", ""), "") if mapping.get("reference") else ""
            prepared.append((index, txn_date, product, quantity, amount, unit_cost, reference))
            if report_type == "purchase":
                supplier_by_row[index] = supplier.strip()
            if report_type == "sales":
                due_by_row[index] = (
                    parse_number(row.get(mapping["cash_due"])) if mapping.get("cash_due") else None,
                    parse_number(row.get(mapping["credit_due"])) if mapping.get("credit_due") else None,
                )
        except (ValueError, TypeError) as error:
            skipped_count += 1
            errors.append("Row {}: {}".format(index, str(error)))
    if not prepared:
        shown = errors[:5]
        suffix = " Aur {} row mein error hai.".format(len(errors) - len(shown)) if len(errors) > len(shown) else ""
        details = "; ".join(shown)
        raise ValueError(
            "Bill Date aur amount ke saath koi valid data row nahi mili. {}{}".format(details, suffix)
        )
    return {
        "filename": filename,
        "content": content,
        "report_type": report_type,
        "prepared": prepared,
        "due_by_row": due_by_row,
        "supplier_by_row": supplier_by_row,
        "skipped_count": skipped_count,
        "skipped_errors": errors,
        "digest": hashlib.sha256(content).hexdigest(),
    }


def backfill_due_data():
    with connect() as connection:
        uploads = connection.execute(
            """SELECT id, filename, source_file FROM uploads
               WHERE report_type = 'sales' AND due_data_loaded = 0"""
        ).fetchall()
        for upload in uploads:
            if upload["source_file"] is None:
                connection.execute(
                    "UPDATE uploads SET due_data_loaded = 1 WHERE id = ?", (upload["id"],)
                )
                continue
            try:
                headers, rows = read_rows(upload["filename"], upload["source_file"])
                mapping = suggested_mapping(headers)
                if not mapping.get("cash_due") and not mapping.get("credit_due"):
                    connection.execute(
                        "UPDATE uploads SET due_data_loaded = 1 WHERE id = ?", (upload["id"],)
                    )
                    continue
                updates = []
                for row_number, row in enumerate(rows, start=2):
                    cash_due = parse_number(row.get(mapping["cash_due"])) if mapping.get("cash_due") else None
                    credit_due = parse_number(row.get(mapping["credit_due"])) if mapping.get("credit_due") else None
                    if cash_due is not None or credit_due is not None:
                        updates.append((cash_due, credit_due, upload["id"], row_number))
                connection.executemany(
                    """UPDATE transactions SET cash_due = ?, credit_due = ?
                       WHERE upload_id = ? AND row_number = ?""",
                    updates,
                )
                connection.execute(
                    "UPDATE uploads SET due_data_loaded = 1 WHERE id = ?", (upload["id"],)
                )
            except (OSError, ValueError) as error:
                print("Could not backfill cash/credit due for {}: {}".format(upload["filename"], error))


def backfill_supplier_data():
    with connect() as connection:
        uploads = connection.execute(
            """SELECT id, filename, source_file FROM uploads
               WHERE report_type = 'purchase' AND supplier_data_loaded = 0"""
        ).fetchall()
        for upload in uploads:
            if upload["source_file"] is None:
                connection.execute(
                    "UPDATE uploads SET supplier_data_loaded = 1 WHERE id = ?", (upload["id"],)
                )
                continue
            try:
                headers, rows = read_rows(upload["filename"], upload["source_file"])
                mapping = suggested_mapping(headers)
                if not mapping.get("supplier"):
                    connection.execute(
                        "UPDATE uploads SET supplier_data_loaded = 1 WHERE id = ?", (upload["id"],)
                    )
                    continue
                updates = [
                    (row.get(mapping["supplier"], "").strip(), upload["id"], row_number)
                    for row_number, row in enumerate(rows, start=2)
                    if row.get(mapping["supplier"], "").strip()
                ]
                connection.executemany(
                    """UPDATE transactions SET supplier = ?
                       WHERE upload_id = ? AND row_number = ?""",
                    updates,
                )
                connection.execute(
                    "UPDATE uploads SET supplier_data_loaded = 1 WHERE id = ?", (upload["id"],)
                )
            except (OSError, ValueError) as error:
                print("Could not backfill supplier data for {}: {}".format(upload["filename"], error))


def aggregate(start, end, hospital, pharmacy):
    conditions = ["t.txn_date >= ?", "t.txn_date <= ?"]
    params = [start, end]
    add_facility_filters(conditions, params, hospital, pharmacy)
    where = " AND ".join(conditions)
    due_conditions = [*conditions, "u.report_type = 'sales'"]
    with connect() as connection:
        totals = connection.execute(
            """SELECT u.report_type, SUM(t.amount) AS amount, COUNT(*) AS rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE """ + where + " GROUP BY u.report_type",
            params,
        ).fetchall()
        facilities = connection.execute(
            "SELECT " + HOSPITAL_SQL + " AS hospital, " + PHARMACY_SQL +
            """ AS pharmacy, u.report_type, SUM(t.amount) AS amount, COUNT(*) AS rows,
                      SUM(t.cash_due) AS cash_due, SUM(t.credit_due) AS credit_due,
                      COUNT(t.cash_due) AS cash_due_rows,
                      COUNT(t.credit_due) AS credit_due_rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE """ + where + " GROUP BY hospital, pharmacy, u.report_type",
            params,
        ).fetchall()
        suppliers = connection.execute(
            "SELECT " + HOSPITAL_SQL + " AS hospital, " + PHARMACY_SQL +
            """ AS pharmacy, TRIM(t.supplier) AS supplier, SUM(t.amount) AS amount
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE """ + where + " AND u.report_type = 'purchase' AND TRIM(t.supplier) != ''"
            " GROUP BY hospital, pharmacy, TRIM(t.supplier) COLLATE NOCASE"
            " ORDER BY amount DESC, supplier COLLATE NOCASE",
            params,
        ).fetchall()
        due = connection.execute(
            """SELECT SUM(t.cash_due) AS cash_due, COUNT(t.cash_due) AS cash_due_rows,
                      SUM(t.credit_due) AS credit_due, COUNT(t.credit_due) AS credit_due_rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE """ + " AND ".join(due_conditions),
            params,
        ).fetchone()
        latest_date = connection.execute(
            """SELECT MAX(t.txn_date) AS latest_date
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE """ + where,
            params,
        ).fetchone()["latest_date"]
    values = {report_type: {"amount": 0, "rows": 0} for report_type in REPORT_TYPES}
    for row in totals:
        values[row["report_type"]] = {"amount": row["amount"] or 0, "rows": row["rows"]}
    by_facility = {}
    for row in facilities:
        key = "{}|{}".format(row["hospital"], row["pharmacy"])
        if key not in by_facility:
            by_facility[key] = {
                "hospital": row["hospital"],
                "pharmacy": row["pharmacy"],
                "sales": 0,
                "collection": 0,
                "purchase": 0,
                "sales_rows": 0,
                "collection_rows": 0,
                "purchase_rows": 0,
                "cash_due": 0,
                "credit_due": 0,
                "cash_due_rows": 0,
                "credit_due_rows": 0,
            }
        by_facility[key][row["report_type"]] = row["amount"] or 0
        by_facility[key][row["report_type"] + "_rows"] = row["rows"]
        if row["report_type"] == "sales":
            by_facility[key]["sales_rows"] = row["rows"]
            by_facility[key]["cash_due"] = row["cash_due"] or 0
            by_facility[key]["credit_due"] = row["credit_due"] or 0
            by_facility[key]["cash_due_rows"] = row["cash_due_rows"] or 0
            by_facility[key]["credit_due_rows"] = row["credit_due_rows"] or 0
    return {
        "latest_date": latest_date,
        "totals": values,
        "due_totals": {
            "cash_due": {"amount": due["cash_due"] or 0, "rows": due["cash_due_rows"]},
            "credit_due": {"amount": due["credit_due"] or 0, "rows": due["credit_due_rows"]},
        },
        "facilities": sorted(by_facility.values(), key=lambda item: (item["hospital"], item["pharmacy"])),
        "supplier_totals": [
            {
                "hospital": row["hospital"],
                "pharmacy": row["pharmacy"],
                "supplier": row["supplier"],
                "amount": row["amount"] or 0,
            }
            for row in suppliers
        ],
    }


def latest_transaction_date():
    with connect() as connection:
        return connection.execute(
            "SELECT MAX(txn_date) AS latest_date FROM transactions"
        ).fetchone()["latest_date"]


def trend(start, end, hospital, pharmacy):
    conditions = ["t.txn_date >= ?", "t.txn_date <= ?"]
    params = [start, end]
    add_facility_filters(conditions, params, hospital, pharmacy)
    with connect() as connection:
        rows = connection.execute(
            """SELECT strftime('%Y-%m', t.txn_date) AS month, u.report_type,
                      SUM(t.amount) AS amount, SUM(t.cash_due) AS cash_due,
                      SUM(t.credit_due) AS credit_due, COUNT(t.cash_due) AS cash_due_rows,
                      COUNT(t.credit_due) AS credit_due_rows, COUNT(*) AS rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE """ + " AND ".join(conditions) +
            " GROUP BY month, u.report_type ORDER BY month",
            params,
        ).fetchall()
    points = {}
    for row in rows:
        point = points.setdefault(
            row["month"],
            {
                "month": row["month"], "sales": 0, "collection": 0, "purchase": 0,
                "cash_due": 0, "credit_due": 0, "cash_due_rows": 0, "credit_due_rows": 0,
                "sales_rows": 0,
            },
        )
        point[row["report_type"]] = row["amount"] or 0
        if row["report_type"] == "sales":
            point["sales_rows"] = row["rows"]
        point["cash_due"] += row["cash_due"] or 0
        point["credit_due"] += row["credit_due"] or 0
        point["cash_due_rows"] += row["cash_due_rows"] or 0
        point["credit_due_rows"] += row["credit_due_rows"] or 0
    return list(points.values())


def hospital_monthly_trend(start, end):
    with connect() as connection:
        rows = connection.execute(
            """SELECT strftime('%Y-%m', t.txn_date) AS month,
                      """ + HOSPITAL_SQL + """ AS hospital,
                      u.report_type, SUM(t.amount) AS amount, COUNT(*) AS rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE t.txn_date >= ? AND t.txn_date <= ?
               GROUP BY month, hospital, u.report_type
               ORDER BY month, hospital""",
            (start, end),
        ).fetchall()
    points = {}
    for row in rows:
        point = points.setdefault(
            (row["month"], row["hospital"]),
            {
                "month": row["month"],
                "hospital": row["hospital"],
                "sales": 0,
                "collection": 0,
                "purchase": 0,
                "sales_rows": 0,
                "collection_rows": 0,
                "purchase_rows": 0,
            },
        )
        point[row["report_type"]] = row["amount"] or 0
        point[row["report_type"] + "_rows"] = row["rows"]
    return list(points.values())


def export_monthly_data():
    with connect() as connection:
        rows = connection.execute(
            """SELECT strftime('%Y-%m', t.txn_date) AS month,
                      """ + HOSPITAL_SQL + """ AS hospital,
                      """ + PHARMACY_SQL + """ AS pharmacy,
                      u.report_type, SUM(t.amount) AS amount,
                      SUM(t.cash_due) AS cash_due, SUM(t.credit_due) AS credit_due,
                      COUNT(t.cash_due) AS cash_due_rows, COUNT(t.credit_due) AS credit_due_rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               GROUP BY month, hospital, pharmacy, u.report_type
               ORDER BY month, hospital, pharmacy"""
        ).fetchall()
    by_facility_month = {}
    for row in rows:
        key = (row["month"], row["hospital"], row["pharmacy"])
        if key not in by_facility_month:
            by_facility_month[key] = {
                "month": row["month"],
                "hospital": row["hospital"],
                "pharmacy": row["pharmacy"],
                "sales": 0,
                "collection": 0,
                "purchase": 0,
                "cash_due": 0,
                "credit_due": 0,
                "cash_due_rows": 0,
                "credit_due_rows": 0,
            }
        point = by_facility_month[key]
        point[row["report_type"]] = row["amount"] or 0
        point["cash_due"] += row["cash_due"] or 0
        point["credit_due"] += row["credit_due"] or 0
        point["cash_due_rows"] += row["cash_due_rows"] or 0
        point["credit_due_rows"] += row["credit_due_rows"] or 0
    return list(by_facility_month.values())


def export_daily_data():
    with connect() as connection:
        rows = connection.execute(
            """SELECT t.txn_date AS date,
                      """ + HOSPITAL_SQL + """ AS hospital,
                      """ + PHARMACY_SQL + """ AS pharmacy,
                      u.report_type, SUM(t.amount) AS amount, COUNT(*) AS rows,
                      SUM(t.cash_due) AS cash_due, SUM(t.credit_due) AS credit_due,
                      COUNT(t.cash_due) AS cash_due_rows, COUNT(t.credit_due) AS credit_due_rows
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               GROUP BY date, hospital, pharmacy, u.report_type
               ORDER BY date, hospital, pharmacy"""
        ).fetchall()
    by_facility_date = {}
    for row in rows:
        key = (row["date"], row["hospital"], row["pharmacy"])
        if key not in by_facility_date:
            by_facility_date[key] = {
                "date": row["date"],
                "hospital": row["hospital"],
                "pharmacy": row["pharmacy"],
                "sales": 0,
                "collection": 0,
                "purchase": 0,
                "sales_rows": 0,
                "collection_rows": 0,
                "purchase_rows": 0,
                "cash_due": 0,
                "credit_due": 0,
                "cash_due_rows": 0,
                "credit_due_rows": 0,
            }
        point = by_facility_date[key]
        point[row["report_type"]] = row["amount"] or 0
        point[row["report_type"] + "_rows"] = row["rows"]
        point["cash_due"] += row["cash_due"] or 0
        point["credit_due"] += row["credit_due"] or 0
        point["cash_due_rows"] += row["cash_due_rows"] or 0
        point["credit_due_rows"] += row["credit_due_rows"] or 0
    return list(by_facility_date.values())


def export_supplier_data():
    with connect() as connection:
        rows = connection.execute(
            """SELECT t.txn_date AS date,
                      """ + HOSPITAL_SQL + """ AS hospital,
                      """ + PHARMACY_SQL + """ AS pharmacy,
                      TRIM(t.supplier) AS supplier, SUM(t.amount) AS amount
               FROM transactions t JOIN uploads u ON u.id = t.upload_id
               WHERE u.report_type = 'purchase' AND TRIM(t.supplier) != ''
               GROUP BY date, hospital, pharmacy, TRIM(t.supplier) COLLATE NOCASE
               ORDER BY date, hospital, pharmacy, supplier COLLATE NOCASE"""
        ).fetchall()
    return [
        {
            "date": row["date"],
            "hospital": row["hospital"],
            "pharmacy": row["pharmacy"],
            "supplier": row["supplier"],
            "amount": row["amount"] or 0,
        }
        for row in rows
    ]


class Handler(BaseHTTPRequestHandler):
    def send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_error_json(self, status, message):
        self.send_json(status, {"error": message})

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/api/health":
            self.send_json(200, {"status": "ok"})
            return
        if path == "/api/export-data":
            self.send_json(200, {
                "hospitals": list(HOSPITALS),
                "latest_date": latest_transaction_date(),
                "pharmacies": {
                    hospital: list(HOSPITAL_PHARMACIES[hospital])
                    for hospital in HOSPITALS
                },
                "monthly": export_monthly_data(),
                "daily": export_daily_data(),
                "suppliers": export_supplier_data(),
            })
            return
        if path == "/api/uploads":
            query = parse_qs(parsed.query)
            hospital = query.get("hospital", ["all"])[0]
            if hospital != "all" and hospital not in HOSPITALS:
                self.send_error_json(400, "Invalid hospital filter.")
                return
            with connect() as connection:
                condition = ""
                params = ()
                if hospital != "all":
                    values = hospital_values(hospital)
                    condition = " WHERE u.hospital IN ({})".format(",".join("?" for _ in values))
                    params = values
                uploads = connection.execute(
                    "SELECT id, " + HOSPITAL_SQL + " AS hospital, " + PHARMACY_SQL +
                    """ AS pharmacy, report_type, filename, row_count, skipped_count, uploaded_at,
                              source_file IS NOT NULL AS source_available
                       FROM uploads u"""
                    + condition + " ORDER BY id DESC LIMIT 100",
                    params,
                ).fetchall()
            self.send_json(200, [dict(upload) for upload in uploads])
            return
        if path.startswith("/api/uploads/") and path.endswith("/file"):
            try:
                upload_id = int(path.split("/")[-2])
            except ValueError:
                self.send_error_json(400, "Invalid upload id.")
                return
            with connect() as connection:
                upload = connection.execute(
                    "SELECT filename, source_file FROM uploads WHERE id = ?",
                    (upload_id,),
                ).fetchone()
            if not upload:
                self.send_error_json(404, "Upload nahi mila.")
                return
            if upload["source_file"] is None:
                self.send_error_json(404, "Original file is not available for this older upload.")
                return
            body = upload["source_file"]
            self.send_response(200)
            self.send_header("Content-Type", "application/octet-stream")
            self.send_header(
                "Content-Disposition",
                "attachment; filename*=UTF-8''{}".format(quote(upload["filename"], safe="")),
            )
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/api/summary":
            query = parse_qs(parsed.query)
            try:
                start = date.fromisoformat(query.get("start", [""])[0]).isoformat()
                end = date.fromisoformat(query.get("end", [""])[0]).isoformat()
                compare_start = date.fromisoformat(query.get("compare_start", [""])[0]).isoformat()
                compare_end = date.fromisoformat(query.get("compare_end", [""])[0]).isoformat()
                if start > end or compare_start > compare_end:
                    raise ValueError("Date range ulta hai.")
                hospital, pharmacy = selected_filters(query)
                hospital_current = aggregate(start, end, "all", "all")
                hospital_comparison = aggregate(compare_start, compare_end, "all", "all")
                self.send_json(200, {
                    "latest_date": latest_transaction_date(),
                    "current": aggregate(start, end, hospital, pharmacy),
                    "comparison": aggregate(compare_start, compare_end, hospital, pharmacy),
                    "hospital_comparison": {
                        "current": {
                            "facilities": hospital_current["facilities"],
                            "supplier_totals": hospital_current["supplier_totals"],
                        },
                        "comparison": {"facilities": hospital_comparison["facilities"]},
                    },
                    "facility_comparison": {
                        "current": aggregate(start, end, hospital, "all")["facilities"],
                        "comparison": aggregate(compare_start, compare_end, hospital, "all")["facilities"],
                    },
                    "trend": trend(start, end, hospital, pharmacy),
                    "trend_comparison": trend(
                        compare_start, compare_end, hospital, pharmacy
                    ),
                    "hospital_monthly": hospital_monthly_trend(start, end),
                })
            except ValueError as error:
                self.send_error_json(400, str(error))
            return
        if path == "/" or path in ("/app.js", "/styles.css"):
            file_path = STATIC / ("index.html" if path == "/" else path.lstrip("/"))
            if not file_path.is_file():
                self.send_error_json(404, "App file nahi mili.")
                return
            body = file_path.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", mimetypes.guess_type(file_path.name)[0] or "text/html; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        self.send_error_json(404, "Route nahi mili.")

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path not in ("/api/preview", "/api/import", "/api/batch-import"):
            self.send_error_json(404, "Route nahi mili.")
            return
        origin = self.headers.get("Origin")
        if origin and origin != "http://127.0.0.1:{}".format(PORT):
            self.send_error_json(403, "Cross-origin request blocked.")
            return
        try:
            files, fields = parse_multipart(self)
            if parsed.path == "/api/preview":
                if len(files) != 1:
                    raise ValueError("Preview one report file at a time.")
                filename, content = files[0]["filename"], files[0]["content"]
                headers, rows = read_rows(filename, content)
                self.send_json(200, {
                    "filename": filename,
                    "headers": headers,
                    "suggestions": suggested_mapping(headers),
                    "row_count": len(rows),
                })
                return
            if parsed.path == "/api/batch-import" and not 1 <= len(files) <= len(REPORT_TYPES):
                raise ValueError("Select between one and three report files.")
            hospital = fields.get("hospital", "")
            if hospital not in HOSPITALS:
                raise ValueError("Select a valid hospital.")
            try:
                if parsed.path == "/api/import":
                    if len(files) != 1:
                        raise ValueError("Import one report file at a time.")
                    entries = [{
                        "pharmacy": fields.get("pharmacy", ""),
                        "report_type": fields.get("report_type", ""),
                        "mapping": json.loads(fields.get("mapping", "{}")),
                    }]
                else:
                    entries = json.loads(fields.get("reports", ""))
            except (json.JSONDecodeError, TypeError) as error:
                raise ValueError("Batch report settings valid nahi hain.") from error
            if not isinstance(entries, list) or len(entries) != len(files):
                raise ValueError("Har selected file ke liye report settings zaroori hain.")
            if parsed.path == "/api/batch-import":
                batch_types = [entry.get("report_type") for entry in entries if isinstance(entry, dict)]
                batch_pharmacies = [entry.get("pharmacy") for entry in entries if isinstance(entry, dict)]
                if len(batch_types) != len(files) or len(set(batch_types)) != len(batch_types):
                    raise ValueError("Upload only one file for each report type at a time.")
                if len(set(batch_pharmacies)) != 1:
                    raise ValueError("Select one pharmacy for all three reports in a batch.")

            prepared_uploads = []
            for file_info, entry in zip(files, entries):
                if not isinstance(entry, dict):
                    raise ValueError("Batch report settings valid nahi hain.")
                pharmacy = entry.get("pharmacy", "")
                if pharmacy not in HOSPITAL_PHARMACIES[hospital]:
                    raise ValueError("Selected pharmacy does not belong to this hospital.")
                report = prepare_import(
                    file_info["filename"],
                    file_info["content"],
                    entry.get("report_type", ""),
                    entry.get("mapping", {}),
                )
                report["pharmacy"] = pharmacy
                prepared_uploads.append(report)

            with connect() as connection:
                result_uploads = []
                for report in prepared_uploads:
                    if duplicate_upload_exists(
                        connection,
                        hospital,
                        report["pharmacy"],
                        report["report_type"],
                        report["digest"],
                    ):
                        raise ValueError(
                            "Yahi file isi hospital, pharmacy aur report type ke liye pehle upload ho chuki hai."
                        )
                    try:
                        cursor = connection.execute(
                            """INSERT INTO uploads
                               (hospital, pharmacy, report_type, filename, sha256, row_count,
                                source_file, skipped_count, due_data_loaded, supplier_data_loaded)
                               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                            (
                                hospital,
                                report["pharmacy"],
                                report["report_type"],
                                report["filename"],
                                report["digest"],
                                len(report["prepared"]),
                                report["content"],
                                report["skipped_count"],
                                1,
                                1,
                            ),
                        )
                    except sqlite3.IntegrityError as error:
                        if "UNIQUE constraint failed" in str(error):
                            raise ValueError(
                                "Yahi file isi hospital, pharmacy aur report type ke liye pehle upload ho chuki hai."
                            ) from error
                        raise
                    connection.executemany(
                        """INSERT INTO transactions
                           (upload_id, row_number, txn_date, product, quantity, amount,
                            unit_cost, reference, cash_due, credit_due, supplier)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        [
                            (
                                cursor.lastrowid, *row,
                                *report["due_by_row"].get(row[0], (None, None)),
                                report["supplier_by_row"].get(row[0], ""),
                            )
                            for row in report["prepared"]
                        ],
                    )
                    result_uploads.append({
                        "filename": report["filename"],
                        "hospital": hospital,
                        "pharmacy": report["pharmacy"],
                        "report_type": REPORT_LABELS[report["report_type"]],
                        "row_count": len(report["prepared"]),
                        "skipped_count": report["skipped_count"],
                    })
            if parsed.path == "/api/import":
                self.send_json(201, {"message": "Data import ho gaya.", **result_uploads[0]})
            else:
                self.send_json(201, {
                    "message": "All files imported and stored.",
                    "hospital": hospital,
                    "file_count": len(result_uploads),
                    "row_count": sum(item["row_count"] for item in result_uploads),
                    "skipped_count": sum(item["skipped_count"] for item in result_uploads),
                    "uploads": result_uploads,
                })
        except ValueError as error:
            self.send_error_json(400, str(error))
        except Exception as error:
            self.log_error("Import failed: %s", error)
            self.send_error_json(500, "Import complete nahi hua. Server log check karein.")

    def do_DELETE(self):
        origin = self.headers.get("Origin")
        if origin and origin != "http://127.0.0.1:{}".format(PORT):
            self.send_error_json(403, "Cross-origin request blocked.")
            return
        path = urlparse(self.path).path
        if not path.startswith("/api/uploads/"):
            self.send_error_json(404, "Route nahi mili.")
            return
        try:
            upload_id = int(path.rsplit("/", 1)[1])
        except ValueError:
            self.send_error_json(400, "Invalid upload id.")
            return
        with connect() as connection:
            deleted = connection.execute("DELETE FROM uploads WHERE id = ?", (upload_id,)).rowcount
        if not deleted:
            self.send_error_json(404, "Upload nahi mila.")
            return
        self.send_json(200, {"message": "Upload aur uske imported rows delete ho gaye."})

    def log_message(self, format_string, *args):
        print("{} - {}".format(self.address_string(), format_string % args))


if __name__ == "__main__":
    init_db()
    print("Pharmacy dashboard is running at http://127.0.0.1:8001")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
